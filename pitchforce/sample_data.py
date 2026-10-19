"""Synthetic sample event (clearly labeled; no real people, companies, or contacts).

The scenario mirrors the structure of real PitchForce recordings so every workflow can be exercised:
- A presenter named differently in the transcript and roster (alias case, needs review).
- A presenting brand that differs from the roster's legal entity with a parent/licensor mention.
- A founder's opening update before the formal introduction and a screen-share interruption.
- An announced total with no individual scores; five 3s + one confirmed-absent judge (18/30);
  an out-of-range score; a missing (blank) score that must not trigger imputation.
- A Zoom multi-section poll CSV with historical launches, multiline multi-select answers, trailing
  blank columns, an anonymous response, and a name/email conflict.
- A wall-clock TXT transcript plus a recording-relative VTT whose recording starts later than the
  meeting and has a paused section (piecewise alignment).
All names use the reserved example.com/example.org domains.
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from openpyxl import Workbook

from .util import fmt_hms

SAMPLE_EVENT = {
    "name": "SAMPLE — Synthetic Life Science Virtual PitchForce",
    "event_date": date(2026, 9, 17),
    "timezone": "America/Los_Angeles",
    "format": "virtual",
    "webinar_id": "812 3456 7890",
}
MEETING_START = 14 * 3600 + 20 * 60 + 33  # 14:20:33 wall clock (meeting transcript begins)
RECORDING_START = 15 * 3600 + 0 * 60 + 14  # 15:00:14 (cloud recording begins later)
PAUSE_AT = 15 * 3600 + 31 * 60  # recording paused here...
PAUSE_LEN = 45.0  # ...for 45 seconds

HOST = "Avery Quinn (EcoVG)"
J = {
    "jordan": "Jordan Avery", "priya": "Priya Nair", "tom": "Tom Becker", "lena": "Lena Ortiz", "sam": "Sam Kowalski",
}
ABSENT_JUDGE = "Morgan Lee"


@dataclass
class Line:
    t: float  # wall-clock seconds
    speaker: str
    text: str


def _pitch(speaker: str, paragraphs: list[str], t: float, step: float = 48.0) -> tuple[list[Line], float]:
    out = []
    for p in paragraphs:
        out.append(Line(t, speaker, p))
        t += step
    return out, t


def build_script() -> list[Line]:
    L: list[Line] = []
    t = float(MEETING_START)
    L.append(Line(t, HOST, "Hi everyone, we'll get started in a few minutes while people join. Please rename yourself with your firm.")); t += 900
    L.append(Line(t, HOST, "Still waiting on a couple of panelists, thanks for your patience.")); t = float(RECORDING_START) + 20
    L.append(Line(t, HOST, "Welcome to the SAMPLE Life Science Virtual PitchForce. Each company has four minutes to pitch, then questions from our investor panel, then feedback. One note: Morgan couldn't join us today, so we have five judges.")); t += 40
    L.append(Line(t, J["jordan"], "Happy to be here, looking forward to the companies today.")); t += 15

    # ---- Company 1: Brightwell Diagnostics (opening update before intro, screen-share interruption)
    L.append(Line(t, "Ana Silva", "Before we start, a quick update for the panel: since we applied we closed our seed extension and enrolled our first forty patients in the feasibility study, so some numbers in the deck are already out of date.")); t += 25
    L.append(Line(t, HOST, "Thanks Ana, great news. Up next we have Brightwell Diagnostics, presented by Ana Silva. Take it away.")); t += 10
    seg, t = _pitch("Ana Silva", [
        "Thank you. Brightwell Diagnostics builds a point of care test for early sepsis. Every hour of delayed treatment raises mortality, and current lab tests take six to twelve hours to come back.",
        "Our cartridge reads three host response markers from a drop of blood in twenty minutes. We have a provisional patent filed on the cartridge chemistry and the reader optics.",
    ], t)
    L += seg
    L.append(Line(t, HOST, "Ana, sorry to interrupt, we lost your screen share. Can you share again?")); t += 8
    L.append(Line(t, "Ana Silva", "Apologies, is that better now? Okay. So the market: there are about one point seven million adult sepsis cases a year in the US and emergency departments are our first customers.")); t += 48
    seg, t = _pitch("Ana Silva", [
        "Our feasibility data so far shows strong agreement with the lab reference in the first forty patients. We plan an FDA 510k submission in late 2027 after the pivotal study.",
        "We are raising three million dollars to fund the pivotal study and manufacturing scale-up, which takes us to submission. Thank you.",
    ], t)
    L += seg
    L.append(Line(t, HOST, "Thank you Ana. Let's open it up for questions from the panel.")); t += 8
    L.append(Line(t, J["priya"], "How many patients do you need in the pivotal study, and what does it cost?")); t += 12
    L.append(Line(t, "Ana Silva", "About four hundred patients across three sites, and we estimate about one point eight million of the raise goes to the study.")); t += 25
    L.append(Line(t, J["tom"], "Who pays for the test, the hospital? Is there a reimbursement code?")); t += 10
    L.append(Line(t, "Ana Silva", "The hospital pays per cartridge. For reimbursement we are still evaluating whether an existing code applies; we don't know yet.")); t += 25
    L.append(Line(t, J["tom"], "And what price per cartridge are you assuming?")); t += 8
    L.append(Line(t, "Ana Silva", "Around ninety dollars per test, which is in line with existing sepsis panels.")); t += 20
    L.append(Line(t, HOST, "Great, now let's hear feedback from the panel.")); t += 6
    L.append(Line(t, J["jordan"], "I love the problem and the speed of the test is compelling. You should tighten the market slide and lead with the clinical outcome.")); t += 22
    L.append(Line(t, J["priya"], "The opening update was strong. I would suggest showing the study timeline next to the spending so we can see how capital converts to milestones.")); t += 22
    L.append(Line(t, J["tom"], "Nothing to add, the others covered it.")); t += 6
    L.append(Line(t, J["lena"], "Clear presentation. Make sure the reimbursement path is explained, because that will decide adoption.")); t += 20
    L.append(Line(t, J["sam"], "Agree with everything said. My advice is to name your first two hospital customers.")); t += 15
    L.append(Line(t, HOST, "Thank you panel. Brightwell's total score is eighteen out of thirty.")); t += 30

    # ---- Company 2: transcript "SightCaption / Leo Park", roster "Owlvision / Sam Whitfield"
    L.append(Line(t, HOST, "Next up is SightCaption with Leo Park. Leo, whenever you're ready.")); t += 8
    seg, t = _pitch("Leo Park", [
        "Thanks. SightCaption makes live captioning glasses for people with hearing loss. Captions appear in the wearer's field of view within half a second.",
        "We have sold eight hundred units through audiology clinics and our monthly revenue is growing twenty percent month over month.",
        "We hold two issued patents on the waveguide display. We are raising two million on a SAFE to expand to retail optical chains.",
        "Our team includes an audiologist co-founder and a former hardware lead from a large consumer electronics company. Thank you.",
    ], t)
    L += seg
    L.append(Line(t, HOST, "Thanks Leo. Questions from the panel?")); t += 5
    L.append(Line(t, J["lena"], "What is the gross margin on the glasses, and how does it change at volume?")); t += 10
    L.append(Line(t, "Leo Park", "Today our gross margin is about thirty five percent and we expect fifty percent at ten thousand units because the display cost drops. Our supplier has given us volume pricing in writing, and that's the main driver of the improvement we model over the next eighteen months.")); t += 40
    L.append(Line(t, J["sam"], "How do you compete with the big platform companies if they add captions to their glasses? And what protects you?")); t += 12
    L.append(Line(t, "Leo Park", "Our patents cover the waveguide, and our clinic channel is hard to replicate. But honestly it depends on how they price it.")); t += 25
    L.append(Line(t, J["jordan"], "What's the exit for investors here, who buys you?")); t += 8
    L.append(Line(t, "Leo Park", "Hearing aid manufacturers are natural acquirers, they have been buying software companies.")); t += 20
    L.append(Line(t, HOST, "Okay, let's move to feedback.")); t += 5
    L.append(Line(t, J["jordan"], "Strong traction. My feedback is to explain the investor return more concretely, the exit horizon and comparables.")); t += 18
    L.append(Line(t, J["lena"], "Good job on the margin answer. I'd recommend a slide on unit economics at volume.")); t += 16
    L.append(Line(t, J["priya"], "Your answers ran long. Try to answer directly and briefly, then stop.")); t += 14
    L.append(Line(t, J["sam"], "I'd encourage you to clarify what protects you from the large platforms beyond the patents.")); t += 15
    L.append(Line(t, J["tom"], "I have nothing to add.")); t += 6
    L.append(Line(t, HOST, "Thank you. SightCaption scored twenty three out of thirty.")); t = max(t + 40, PAUSE_AT + PAUSE_LEN + 10)

    # ---- Company 3: transcript "Petrix" (product/licensee) vs roster "Corvane Pharma"
    L.append(Line(t, HOST, "Welcome back after our short break. Next we have Petrix, presented by Ravi Mehta.")); t += 8
    seg, t = _pitch("Ravi Mehta", [
        "Thank you. Petrix is an oral therapy for chronic kidney disease in cats. We hold an exclusive license from Corvain Pharma, our parent company, for animal health uses of the molecule.",
        "In a pilot study the drug reduced a key kidney marker in twelve of fifteen cats. We have a grant from a veterinary foundation and expect conditional approval in 2028.",
        "The market for feline kidney disease is large because about a third of older cats are affected. We are raising four million to complete the pivotal field study.",
    ], t)
    L += seg
    L.append(Line(t, HOST, "Thank you Ravi. Let's go to questions.")); t += 5
    L.append(Line(t, J["tom"], "What is the regulatory pathway and how long is the pivotal study?")); t += 10
    L.append(Line(t, "Ravi Mehta", "It's the conditional approval pathway with the veterinary regulator. The pivotal field study is about eighteen months [inaudible]")); t += 20
    L.append(Line(t, J["priya"], "How does the license from the parent work if Petrix is acquired, does it")); t += 6
    L.append(Line(t, "Ravi Mehta", "That's a good question, we're still working on that with our lawyers.")); t += 15
    L.append(Line(t, HOST, "Thank you. Feedback from the panel please.")); t += 5
    L.append(Line(t, J["tom"], "Promising data. You should simplify the slides, they were dense and hard to read.")); t += 15
    L.append(Line(t, J["priya"], "I'd recommend clarifying the corporate structure and the license terms early in the pitch.")); t += 15
    L.append(Line(t, J["lena"], "Nothing more from me.")); t += 5
    L.append(Line(t, HOST, "Thanks everyone. Petrix received a total of twenty one out of thirty.")); t += 30

    # ---- Company 4: Hemaflow Bio (announced total only; five panelists speak)
    L.append(Line(t, HOST, "Our last company is Hemaflow Bio with Dana Okafor. Please welcome Dana.")); t += 8
    seg, t = _pitch("Dana Okafor", [
        "Thank you. Hemaflow Bio is developing a shelf stable blood substitute for trauma care in the field.",
        "We have completed formulation work and we're preparing our first animal study. We submitted a proposal to a federal biodefense program and are waiting to hear back.",
        "Our cell bank build will take about nine months. We are raising five million to complete the animal studies and the cell bank.",
    ], t)
    L += seg
    L.append(Line(t, HOST, "Thank you Dana. Questions?")); t += 5
    L.append(Line(t, J["jordan"], "Where are you in development and what is the status of the federal submission?")); t += 10
    L.append(Line(t, "Dana Okafor", "We finished formulation and the submission is under review; we expect a decision this winter.")); t += 20
    L.append(Line(t, J["lena"], "Which animal model will the regulator accept, do you have access to it, and who are your competitors?")); t += 12
    L.append(Line(t, "Dana Okafor", "We're planning a swine model with a contract research organization, and the two main competitors are academic programs.")); t += 22
    L.append(Line(t, J["priya"], "When does the animal study start and finish?")); t += 8
    L.append(Line(t, "Dana Okafor", "It starts next quarter and runs about six months.")); t += 15
    L.append(Line(t, HOST, "Let's hear the panel's feedback.")); t += 5
    L.append(Line(t, J["priya"], "Connect the animal study timeline to your spending so we can see what the five million buys.")); t += 15
    L.append(Line(t, J["jordan"], "Shorten your answers; lead with the direct answer and stop.")); t += 12
    L.append(Line(t, J["sam"], "You should clarify your IP and competitive protection, and the timelines to get there.")); t += 14
    L.append(Line(t, J["tom"], "Explain the investor return and the liquidity horizon; that's what we need to underwrite.")); t += 14
    L.append(Line(t, J["lena"], "The slides were dense. Simplify them so a non-specialist can follow.")); t += 14
    L.append(Line(t, HOST, "Thank you. Hemaflow's total is twenty one out of thirty.")); t += 60
    L.append(Line(t, HOST, "And our runner-up today is Hemaflow Bio, congratulations to Dana and the team.")); t += 10
    L.append(Line(t, HOST, "Our winner is SightCaption. Thank you all for joining.")); t += 10
    return L


def write_txt(lines: list[Line]) -> str:
    out = io.StringIO()
    for ln in lines:
        out.write(f"[{ln.speaker}] {fmt_hms(ln.t)}\n{ln.text}\n\n")
    return out.getvalue()


def media_time(wall: float) -> float | None:
    """Wall clock -> media time for the synthetic recording (starts late, pauses once)."""
    if wall < RECORDING_START:
        return None
    if PAUSE_AT <= wall < PAUSE_AT + PAUSE_LEN:
        return None
    m = wall - RECORDING_START
    if wall >= PAUSE_AT + PAUSE_LEN:
        m -= PAUSE_LEN
    return m


def write_vtt(lines: list[Line]) -> str:
    out = io.StringIO()
    out.write("WEBVTT\n\n")
    n = 0
    for i, ln in enumerate(lines):
        m = media_time(ln.t)
        if m is None:
            continue
        nxt = lines[i + 1].t if i + 1 < len(lines) else ln.t + 8
        m_end = m + max(1.0, min(nxt - ln.t, 60) - 0.2)
        n += 1
        out.write(f"{n}\n{fmt_hms(m, millis=True)} --> {fmt_hms(m_end, millis=True)}\n{ln.speaker}: {ln.text}\n\n")
    return out.getvalue()


def recording_duration(lines: list[Line]) -> float:
    return media_time(lines[-1].t) + 15


COMPANY_ROSTER = [
    ["Order", "Company", "Presenter", "Email", "Phone", "Website"],
    ["1", "Brightwell Diagnostics", "Ana Silva", "ana@brightwell.example.com", "", "https://brightwell.example.com"],
    ["2", "Owlvision", "Sam Whitfield", "sam@owlvision.example.com", "+1 555 0100", ""],
    ["3", "Corvane Pharma", "Ravi Mehta", "ravi@corvane.example.com", "", ""],
    ["4", "Hemaflow Bio", "Dana Okafor", "dana@hemaflow.example.com", "", ""],
]
JUDGE_ROSTER = [
    ["Name", "Affiliation", "Email", "Phone", "Attendance"],
    [J["jordan"], "Northbay Angels (sample)", "jordan@northbay.example.org", "", "Present"],
    [J["priya"], "Cedar Health Fund (sample)", "", "", "Present"],
    [J["tom"], "Becker Family Office (sample)", "tom.becker@example.org", "", "Present"],
    [J["lena"], "", "lena@ortiz.example.org", "", "Present"],
    [J["sam"], "Kowalski Ventures (sample)", "", "", ""],
    [ABSENT_JUDGE, "Lee Capital (sample)", "morgan@leecap.example.org", "", "Absent"],
]


def write_csv(rows: list[list[str]]) -> bytes:
    buf = io.StringIO()
    csv.writer(buf, lineterminator="\r\n").writerows(rows)
    return buf.getvalue().encode("utf-8")


def scores_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Scores"
    ws.append(["Company", J["jordan"], J["priya"], J["tom"], J["lena"], J["sam"], ABSENT_JUDGE])
    ws.append(["Brightwell Diagnostics", 3, 3, 3, 3, 3, None])  # five 3s + confirmed absent judge -> 18/30
    ws.append(["Owlvision", 4, 4.5, 7, 4, 4, None])  # 7 is out of range -> flagged
    ws.append(["Corvane Pharma", 4, 3.5, None, 4, 4, None])  # Tom blank = missing data, no imputation
    # Hemaflow intentionally has no individual scores (announced total only)
    ws2 = wb.create_sheet("Notes")
    ws2.append(["Synthetic sample data. Not real scores."])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def poll_csv_hemaflow() -> bytes:
    """Zoom-style multi-section poll report with historical launches."""
    rows: list[list[str]] = []
    rows.append(["Poll Report", "", "", "", "", "", ""])
    rows.append(["Report Generated:", "2026-09-18 09:00:00", "", "", "", "", ""])
    rows.append(["Overview", "", "", "", "", "", ""])
    rows.append(["Poll Name", "Type", "Created by", "Launches", "", "", ""])
    rows.append(["Follow-Up- Presenting Hemaflow Bio", "Multiple choice", "SAMPLE organizer", "4", "", "", ""])
    rows.append(["", "", "", "", "", "", ""])
    rows.append(["Launch History", "", "", "", "", "", ""])
    rows.append(["#", "Meeting/Webinar ID", "Topic", "Launch Start Time", "Launch End Time", "Participants", ""])
    rows.append(["1", "812 0000 1111", "Spring PitchForce (sample)", "2026-03-12 16:05:00", "2026-03-12 16:09:00", "40", ""])
    rows.append(["2", "812 0000 2222", "Summer PitchForce (sample)", "2026-06-18 15:55:10", "2026-06-18 15:58:00", "35", ""])
    rows.append(["3", "812 3456 7890", "SAMPLE Life Science Virtual PitchForce", "2026-09-17 15:58:40", "2026-09-17 16:03:00", "52", ""])
    rows.append(["4", "812 0000 3333", "Autumn mixer (sample)", "2026-09-24 12:00:00", "2026-09-24 12:04:00", "20", ""])
    rows.append(["", "", "", "", "", "", ""])
    rows.append(["Poll Details", "", "", "", "", "", ""])
    q = "What follow-up would you like with this company?"
    rows.append(["#", "User Name", "Email Address", "Submitted Date and Time", q, "", ""])
    hist = [
        ("1", "Chris Lane", "chris@lane.example.org", "2026-03-12 16:06:10", "I'm interested in investing.\nI'd like a meeting."),
        ("2", "Jordan Avery", "jordan.old@example.org", "2026-03-12 16:06:40", "I need more info."),
        ("3", "Kim Patel", "kim@patel.example.org", "2026-06-18 15:56:00", "I have a free resource."),
        ("4", "Priya Nair", "priya.n@example.net", "2026-06-18 15:56:30", "I'm interested in investing."),
    ]
    event_rows = [
        ("49", "Pat Rivera", "pat.rivera@example.org", "2026-09-17 16:00:44", "I need more info."),
        ("50", "Alex Marin", "tom.becker@example.org", "2026-09-17 16:00:51", "I’m not interested."),
        ("51", "Robin Shaw", "robin@shaw.example.org", "2026-09-17 16:01:02", "I'm not interested."),
        ("52", "", "", "2026-09-17 16:01:15", "I'm not interested."),
        ("53", "Casey Wu", "casey@wu.example.org", "2026-09-17 16:01:30", "I'm not interested."),
    ]
    later = [("54", "Drew Hall", "drew@hall.example.org", "2026-09-24 12:01:00", "I'd like a meeting.")]
    for r in hist + event_rows + later:
        rows.append(list(r) + ["", ""])
    return write_csv(rows)


def interest_xlsx_other() -> bytes:
    """Ordinary workbook with a Company column: Brightwell + Owlvision responses (two sheets)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Brightwell"
    ws.append(["Name", "Email", "Submitted", "Company", "Response"])
    ws.append(["Jamie Fox", "jamie@fox.example.org", "2026-09-17 15:24:00", "Brightwell Diagnostics", "I'm interested in investing.;I'd like a meeting."])
    ws.append(["Jamie Fox", "jamie@fox.example.org", "2026-09-17 15:26:30", "Brightwell Diagnostics", "I'm interested in investing.;I'd like a meeting.;I can make an introduction."])
    ws.append(["Anonymous", "", "2026-09-17 15:25:00", "Brightwell Diagnostics", "I need more info."])
    ws.append(["Lee Grant", "lee@grant.example.org", "2026-09-17 15:25:40", "Brightwell Diagnostics", "I'm not interested.;I'd like a meeting."])
    ws.append(["Morgan Diaz", "morgan@diaz.example.org", "2026-06-18 15:30:00", "Brightwell Diagnostics", "I'm interested in investing."])
    ws2 = wb.create_sheet("Owlvision")
    ws2.append(["Name", "Email", "Submitted", "Company", "Response"])
    ws2.append(["Taylor Brooks", "taylor@brooks.example.org", "2026-09-17 15:38:00", "SightCaption", "I'm open to mentoring."])
    ws2.append(["=HYPERLINK(\"http://evil.example\",\"x\")", "x@evil.example.org", "2026-09-17 15:38:20", "SightCaption", "I have a free resource."])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate(dest: Path, with_video: bool = True) -> dict[str, Path]:
    """Write all sample files into ``dest`` and return {kind_label: path}."""
    dest.mkdir(parents=True, exist_ok=True)
    lines = build_script()
    files = {
        "transcript_txt": dest / "SAMPLE_meeting_transcript.txt",
        "transcript_vtt": dest / "SAMPLE_recording_transcript.vtt",
        "company_roster": dest / "SAMPLE_company_roster.csv",
        "judge_roster": dest / "SAMPLE_judge_roster.csv",
        "scores": dest / "SAMPLE_judge_scores.xlsx",
        "interest_csv": dest / "report_Follow-Up- Presenting Hemaflow Bio (SAMPLE).csv",
        "interest_xlsx": dest / "SAMPLE_interest_responses.xlsx",
    }
    files["transcript_txt"].write_text(write_txt(lines), encoding="utf-8")
    files["transcript_vtt"].write_text(write_vtt(lines), encoding="utf-8")
    files["company_roster"].write_bytes(write_csv(COMPANY_ROSTER))
    files["judge_roster"].write_bytes(write_csv(JUDGE_ROSTER))
    files["scores"].write_bytes(scores_xlsx())
    files["interest_csv"].write_bytes(poll_csv_hemaflow())
    files["interest_xlsx"].write_bytes(interest_xlsx_other())
    if with_video:
        from .media.ffmpeg import make_test_video

        files["video"] = dest / "SAMPLE_recording.mp4"
        if not files["video"].exists():
            make_test_video(files["video"], recording_duration(lines), size="160x90", rate=5)
    return files
