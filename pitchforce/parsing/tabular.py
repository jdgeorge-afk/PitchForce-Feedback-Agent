"""Reading CSV and XLSX into plain rows while keeping source references (sheet, row, cell)."""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ..util import norm_name, normalize_punct


@dataclass
class Sheet:
    name: str
    rows: list[list[str]]
    row_refs: list[int] = field(default_factory=list)  # 1-based row number (XLSX) or record number (CSV)
    line_refs: list[int] = field(default_factory=list)  # CSV physical line where the record ends
    kind: str = "csv"

    def cell_ref(self, row_idx: int, col_idx: int) -> str:
        if self.kind == "xlsx":
            return f"{self.name}!{get_column_letter(col_idx + 1)}{self.row_refs[row_idx]}"
        return f"record {self.row_refs[row_idx]}, column {col_idx + 1}"


def _cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16"):
        try:
            text = data.decode(enc)
            if enc == "utf-16" and not data[:2] in (b"\xff\xfe", b"\xfe\xff"):
                continue
            return text
        except UnicodeDecodeError:
            continue
    return data.decode("cp1252", errors="replace")


def read_csv_bytes(data: bytes, name: str = "CSV") -> Sheet:
    text = decode_text(data)
    sample = text[:4096]
    delim = ","
    try:
        delim = csv.Sniffer().sniff(sample, delimiters=",;\t").delimiter
    except csv.Error:
        pass
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delim)
    rows, refs, lines = [], [], []
    for rec_no, row in enumerate(reader, start=1):
        rows.append([c for c in row])
        refs.append(rec_no)
        lines.append(reader.line_num)
    return Sheet(name=name, rows=rows, row_refs=refs, line_refs=lines, kind="csv")


def read_xlsx_bytes(data: bytes) -> list[Sheet]:
    """Cell values as text. Formula cells without a cached value keep their formula text as a plain
    string (never evaluated), so uploaded content is preserved and later neutralized on export."""
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    wf = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    sheets = []
    for ws, wsf in zip(wb.worksheets, wf.worksheets):
        rows, refs = [], []
        for r_idx, (row, frow) in enumerate(zip(ws.iter_rows(values_only=True), wsf.iter_rows(values_only=True)), start=1):
            vals = []
            for v, fv in zip(row, frow):
                if v is None and isinstance(fv, str) and fv.startswith("="):
                    v = fv
                vals.append(_cell_str(v))
            if any(v.strip() for v in vals):
                rows.append(vals)
                refs.append(r_idx)
        sheets.append(Sheet(name=ws.title, rows=rows, row_refs=refs, kind="xlsx"))
    wb.close()
    wf.close()
    return sheets


def read_tabular(data: bytes, filename: str) -> list[Sheet]:
    ext = Path(filename).suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        return read_xlsx_bytes(data)
    if ext == ".xls":
        raise ValueError("Legacy .xls files are not supported. Save the workbook as .xlsx or CSV and upload again.")
    return [read_csv_bytes(data, name=Path(filename).stem[:40] or "CSV")]


def header_key(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", norm_name(normalize_punct(s))).strip()


def find_header(sheet: Sheet, aliases: dict[str, list[str]], min_hits: int = 2, start: int = 0, max_scan: int = 40):
    """Find the first row that looks like a header. Returns (row_index, {field: col_index})."""
    for i in range(start, min(len(sheet.rows), start + max_scan)):
        mapping = map_columns(sheet.rows[i], aliases)
        if len(mapping) >= min_hits:
            return i, mapping
    return None, {}


def map_columns(header: list[str], aliases: dict[str, list[str]]) -> dict[str, int]:
    keys = [header_key(h) if h.strip() != "#" else "#" for h in header]
    mapping: dict[str, int] = {}
    # Pass 1: exact header matches; pass 2: prefix matches for remaining fields.
    for exact in (True, False):
        for fld, names in aliases.items():
            if fld in mapping:
                continue
            for j, k in enumerate(keys):
                if not k or j in mapping.values():
                    continue
                if any(k == a for a in names) if exact else any(len(a) > 3 and k.startswith(a) for a in names):
                    mapping[fld] = j
                    break
    return mapping


def get(row: list[str], mapping: dict[str, int], fld: str) -> str:
    j = mapping.get(fld)
    if j is None or j >= len(row):
        return ""
    return (row[j] or "").strip()
