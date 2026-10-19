"""Excel exports.

- Founder workbook: only this company's event-specific responses. Actionable follow-up sheet,
  original responses for traceability, and a compact summary. No other company's contacts.
- Organizer audit workbook: full event detail (scores, all poll rows incl. historical, launches,
  review items, sources). Never included in founder packages.
All text cells pass through ``excel_safe`` to neutralize formula injection; dates and numbers are typed.
"""
from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import (
    INTEREST_LABELS,
    Company,
    ContactMatch,
    Event,
    Person,
    PollLaunch,
    PollResponse,
    ReviewItem,
    Score,
    SourceFile,
)
from ..services import interest as interest_svc
from ..services.scoring import current_result, rank_companies
from ..util import excel_safe

HEADER_FILL = PatternFill("solid", fgColor="1F4E5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")
OUTREACH_STATUSES = ["Not started", "Contacted", "Replied", "Meeting scheduled", "Materials sent", "Closed – not a fit", "Do not contact"]
CAT_ORDER = ["invest", "meeting", "due_diligence", "more_info", "introduction", "mentoring", "free_resource", "unknown", "not_interested"]
LEAD_TYPE = {"invest": "Investment lead", "meeting": "Meeting request", "due_diligence": "Due-diligence request",
             "more_info": "Information request", "introduction": "Introduction offer", "mentoring": "Mentoring offer",
             "free_resource": "Resource offer", "unknown": "Unclassified response", "not_interested": "Not interested"}


def _sheet(wb: Workbook, title: str, headers: list[str], widths: list[int], first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title[:31]
    ws.append(headers)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(row=1, column=i)
        c.fill, c.font, c.alignment = HEADER_FILL, HEADER_FONT, Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    return ws


def _row(ws, values: list, wrap_cols: set[int] | None = None, date_cols: set[int] | None = None):
    ws.append([excel_safe(v) for v in values])
    r = ws.max_row
    for i in range(1, len(values) + 1):
        cell = ws.cell(row=r, column=i)
        if wrap_cols and i in wrap_cols:
            cell.alignment = WRAP
        if date_cols and i in date_cols and isinstance(values[i - 1], datetime):
            cell.number_format = "yyyy-mm-dd hh:mm:ss"


def _finish(ws):
    if ws.max_row > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"


def _cats(r: PollResponse) -> list[str]:
    cs = sorted({c.category for c in r.choices} or {"unknown"}, key=CAT_ORDER.index)
    return cs


def founder_workbook(db: Session, event: Event, company: Company) -> tuple[bytes, dict]:
    summary = interest_svc.summarize(db, event, company)
    wb = Workbook()
    ws = _sheet(wb, "Follow-up", [
        "Lead type", "Name (as supplied)", "Email (as supplied)", "Phone (if supplied)", "Affiliation (if supplied)",
        "Exact response", "Interest category", "Requested information", "Suggested next action", "Outreach status", "Notes",
        "Submitted", "Source reference", "Identity / data note"],
        [20, 24, 30, 16, 22, 40, 26, 30, 44, 18, 30, 19, 34, 40], first=True)
    rows = [r for r in summary.responses if not r.superseded_by_id and not r.duplicate_of_id]
    actionable = []
    for r in rows:
        cats = _cats(r)
        if set(cats) == {"not_interested"}:
            continue  # explicit not-interested: preserved in the responses sheet only
        actionable.append((CAT_ORDER.index(cats[0]), r, cats))
    actionable.sort(key=lambda x: (x[0], x[1].submitted_at or datetime.min))
    for _, r, cats in actionable:
        notes = []
        if r.is_anonymous:
            notes.append("Anonymous response: no contact details were supplied.")
        if r.has_conflict:
            notes.append("Contradictory selections (includes “not interested”); confirm before outreach.")
        for cm in db.scalars(select(ContactMatch).where(ContactMatch.response_id == r.id, ContactMatch.status == "conflict")):
            notes.append("The email supplied is assigned to a different person on the event roster; identity not confirmed.")
        if r.email_raw and not r.email:
            notes.append("Email as supplied does not look valid.")
        primary = [c for c in cats if c != "not_interested"][0] if [c for c in cats if c != "not_interested"] else cats[0]
        requested = r.answer_raw if any(c in ("more_info", "due_diligence") for c in cats) else ""
        _row(ws, [
            LEAD_TYPE[primary], r.name or "(anonymous)", r.email_raw, r.source_ref.get("phone", ""), r.source_ref.get("company_affiliation", ""),
            r.answer_raw, "; ".join(INTEREST_LABELS[c] for c in cats), requested,
            " ".join(dict.fromkeys(interest_svc.SUGGESTED_ACTION.get(c, "") for c in cats if c != "not_interested")),
            r.outreach_status or "Not started", r.outreach_note, r.submitted_at or r.submitted_raw,
            r.source_ref.get("ref", ""), " ".join(notes)],
            wrap_cols={6, 8, 9, 11, 13, 14}, date_cols={12})
    dv = DataValidation(type="list", formula1='"' + ",".join(OUTREACH_STATUSES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    if ws.max_row > 1:
        dv.add(f"J2:J{ws.max_row}")
    if ws.max_row == 1:
        ws.append([excel_safe("No actionable follow-up responses for this company." if summary.status == "available" else summary.note)])
    _finish(ws)

    ws2 = _sheet(wb, "Event responses", [
        "Original response ID", "Name (as supplied)", "Email (as supplied)", "Submitted (as recorded)", "Submitted",
        "Question", "Response (original text)", "Choice", "Category", "Status", "Source file", "Source reference"],
        [12, 24, 30, 20, 19, 36, 40, 34, 26, 30, 34, 34])
    for r in summary.responses:
        status = "Superseded by a later submission" if r.superseded_by_id else ("Duplicate" if r.duplicate_of_id else "Current")
        if r.is_anonymous:
            status += " · anonymous"
        src = db.get(SourceFile, r.source_file_id)
        choices = r.choices or []
        if not choices:
            _row(ws2, [r.original_response_id, r.name_raw, r.email_raw, r.submitted_raw, r.submitted_at, r.question_raw, r.answer_raw,
                       "", INTEREST_LABELS["unknown"], status, src.original_filename if src else "", r.source_ref.get("ref", "")],
                 wrap_cols={6, 7}, date_cols={5})
        for c in choices:
            _row(ws2, [r.original_response_id, r.name_raw, r.email_raw, r.submitted_raw, r.submitted_at, r.question_raw, r.answer_raw,
                       c.choice_raw, INTEREST_LABELS[c.category], status, src.original_filename if src else "", r.source_ref.get("ref", "")],
                 wrap_cols={6, 7, 8}, date_cols={5})
    _finish(ws2)

    ws3 = _sheet(wb, "Summary", ["Category", "Unique respondents"], [36, 20])
    if summary.status == "available":
        for cat, _ in [(c, None) for c in CAT_ORDER]:
            if cat in summary.by_category:
                _row(ws3, [INTEREST_LABELS[cat], summary.by_category[cat]])
        _row(ws3, ["Total unique respondents", summary.respondents])
        _row(ws3, ["  of which anonymous", summary.anonymous])
        ws3.append([])
        ws3.append([excel_safe("A respondent who chose several options is counted once in each category and once in the total.")])
        ws3.append([excel_safe("“Needs more information” is an information request, not investment interest.")])
    else:
        ws3.append([excel_safe(summary.note)])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), {"status": summary.status, "actionable": len(actionable), "responses": len(summary.responses)}


def audit_workbook(db: Session, event: Event) -> bytes:
    wb = Workbook()
    companies = list(db.scalars(select(Company).where(Company.event_id == event.id).order_by(Company.slot_order)))
    ws = _sheet(wb, "Results", ["Company", "Score status", "Expected judges", "Verified scores", "Raw sum", "Mean", "Confirmed missing",
                                "Imputed contribution", "Adjusted total", "Denominator", "Announced total", "Discrepancy",
                                "Rank (ties shared)", "Explanation"], [28, 22, 10, 10, 10, 10, 10, 12, 12, 10, 12, 11, 12, 70], first=True)
    results = {c.id: current_result(db, c.id) for c in companies}
    ranking = {cid: (rank, tied) for cid, v, rank, tied in rank_companies(
        [(c.id, (results[c.id].adjusted_total if results[c.id] and results[c.id].status in ("verified", "verified_with_imputation") else None))
         for c in companies])}
    for c in companies:
        r = results[c.id]
        rank, tied = ranking.get(c.id, (None, False))
        if r is None:
            _row(ws, [c.display_name, "not computed"])
            continue
        _row(ws, [c.display_name, r.status, r.expected_judges, r.valid_count, r.raw_sum, r.mean, r.confirmed_missing,
                  r.imputed_contribution, r.adjusted_total, r.denominator, r.announced_value, r.discrepancy,
                  (f"{rank} (tie)" if tied else rank) if rank else "", " ".join((r.details or {}).get("explanation", []))],
             wrap_cols={14})
    _finish(ws)
    ws = _sheet(wb, "Judge scores", ["Company (as in source)", "Matched company", "Judge (as in source)", "Matched judge", "Value (as in source)",
                                     "Value", "Status", "Reason", "Source", "Reference"], [26, 26, 24, 24, 12, 8, 18, 50, 26, 26])
    for s in db.scalars(select(Score).where(Score.event_id == event.id)):
        c = db.get(Company, s.company_id) if s.company_id else None
        p = db.get(Person, s.judge_person_id) if s.judge_person_id else None
        src = db.get(SourceFile, s.source_file_id) if s.source_file_id else None
        _row(ws, [s.company_label_raw, c.display_name if c else "", s.judge_label_raw, p.display_name if p else "", s.value_raw, s.value,
                  s.status, s.status_reason, src.original_filename if src else s.source_kind, (s.source_ref or {}).get("ref", "")], wrap_cols={8})
    _finish(ws)
    ws = _sheet(wb, "Poll launches", ["File", "Launch #", "Poll name", "Meeting/Webinar ID", "Start (as recorded)", "Start", "In this event",
                                      "Reason", "Company", "Company status"], [30, 8, 34, 18, 20, 19, 10, 50, 24, 14])
    for pl in db.scalars(select(PollLaunch).where(PollLaunch.event_id == event.id)):
        src = db.get(SourceFile, pl.source_file_id)
        c = db.get(Company, pl.company_id) if pl.company_id else None
        _row(ws, [src.original_filename if src else "", pl.launch_index, pl.poll_name, pl.webinar_id_raw, pl.start_raw, pl.start_at,
                  "yes" if pl.in_event else "no", pl.in_event_reason, c.display_name if c else "", pl.company_status],
             wrap_cols={8}, date_cols={6})
    _finish(ws)
    ws = _sheet(wb, "All poll responses", ["File", "Original ID", "Name", "Email", "Submitted", "Response", "Categories", "In this event",
                                           "Company", "Superseded/duplicate", "Contradictory", "Review note", "Reference"],
                [28, 10, 22, 28, 19, 40, 30, 10, 22, 16, 12, 40, 30])
    for r in db.scalars(select(PollResponse).where(PollResponse.event_id == event.id)):
        src = db.get(SourceFile, r.source_file_id)
        c = db.get(Company, r.company_id) if r.company_id else None
        _row(ws, [src.original_filename if src else "", r.original_response_id, r.name_raw, r.email_raw, r.submitted_at or r.submitted_raw,
                  r.answer_raw, "; ".join(INTEREST_LABELS[x] for x in _cats(r)), "yes" if r.in_event else "no",
                  c.display_name if c else "", "superseded" if r.superseded_by_id else ("duplicate" if r.duplicate_of_id else ""),
                  "yes" if r.has_conflict else "", r.review_reason, r.source_ref.get("ref", "")], wrap_cols={6, 12}, date_cols={5})
    _finish(ws)
    ws = _sheet(wb, "Review items", ["Status", "Severity", "Type", "Company", "Item", "Detail", "Created", "Resolved"],
                [10, 9, 18, 22, 50, 70, 19, 19])
    for it in db.scalars(select(ReviewItem).where(ReviewItem.event_id == event.id).order_by(ReviewItem.created_at)):
        c = db.get(Company, it.company_id) if it.company_id else None
        _row(ws, [it.status, it.severity, it.kind, c.display_name if c else "", it.title, it.detail, it.created_at, it.resolved_at],
             wrap_cols={5, 6}, date_cols={7, 8})
    _finish(ws)
    ws = _sheet(wb, "Sources", ["File", "Type", "Status", "Size (bytes)", "SHA-256", "Uploaded", "Summary"], [40, 18, 10, 12, 66, 19, 60])
    for s in db.scalars(select(SourceFile).where(SourceFile.event_id == event.id)):
        summ = s.parse_summary or {}
        brief = ", ".join(f"{k}={v}" for k, v in summ.items() if not isinstance(v, (list, dict)))[:300]
        _row(ws, [s.original_filename, s.kind, s.status, s.size, s.sha256, s.created_at, brief], wrap_cols={7}, date_cols={6})
    _finish(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
