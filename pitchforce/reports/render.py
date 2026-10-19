"""PDF (ReportLab) and DOCX (python-docx) rendering of a founder report."""
from __future__ import annotations

import html
import io
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

from ..models import ReportSection
from ..util import clean_text
from . import markup

THEME = {"primary": "#1F4E5F", "accent": "#C8702A", "text": "#222222", "muted": "#666666"}


def _rl_spans(spans: list[markup.Span]) -> str:
    out = []
    for s in spans:
        t = html.escape(clean_text(s.text)).replace("\n", "<br/>")
        if s.kind == "bold":
            out.append(f"<b>{t}</b>")
        elif s.kind == "link":
            out.append(f'<link href="{html.escape(s.url)}" color="{THEME["primary"]}"><u>{t}</u></link>')
        else:
            out.append(t)
    return "".join(out)


def render_pdf(title: str, subtitle: str, sections: list[ReportSection], footer: str, theme: dict | None = None) -> bytes:
    th = {**THEME, **(theme or {})}
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=LETTER, leftMargin=0.85 * inch, rightMargin=0.85 * inch, topMargin=0.8 * inch,
                            bottomMargin=0.8 * inch, title=title, author="EcoVG PitchForce", subject=subtitle)
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["BodyText"], fontName="Helvetica", fontSize=10.2, leading=14.2,
                          textColor=colors.HexColor(th["text"]), spaceAfter=6, alignment=TA_LEFT)
    h1 = ParagraphStyle("h1", parent=body, fontName="Helvetica-Bold", fontSize=18, leading=22, textColor=colors.HexColor(th["primary"]), spaceAfter=2)
    sub = ParagraphStyle("sub", parent=body, fontSize=10, textColor=colors.HexColor(th["muted"]), spaceAfter=12)
    h2 = ParagraphStyle("h2", parent=body, fontName="Helvetica-Bold", fontSize=12.5, leading=16, textColor=colors.HexColor(th["primary"]),
                        spaceBefore=10, spaceAfter=4, keepWithNext=1)
    h3 = ParagraphStyle("h3", parent=body, fontName="Helvetica-Bold", fontSize=10.8, leading=14, textColor=colors.HexColor(th["accent"]),
                        spaceBefore=6, spaceAfter=3, keepWithNext=1)
    note = ParagraphStyle("note", parent=body, fontName="Helvetica-Oblique", textColor=colors.HexColor(th["muted"]))
    story = [Paragraph(html.escape(title), h1), Paragraph(html.escape(subtitle), sub)]
    for sec in sections:
        text = sec.text
        if not text.strip():
            continue
        if sec.key != "header":
            story.append(Paragraph(html.escape(sec.title), h2))
        for b in markup.parse(text):
            if b.kind == "heading":
                story.append(Paragraph(_rl_spans(b.spans), h3))
            elif b.kind == "bullets":
                story.append(ListFlowable([ListItem(Paragraph(_rl_spans(i), body), leftIndent=12) for i in b.items],
                                          bulletType="bullet", start="•", leftIndent=12, bulletFontSize=8))
            elif b.kind == "note":
                story.append(Paragraph(_rl_spans(b.spans), note))
            else:
                story.append(Paragraph(_rl_spans(b.spans), body))
        story.append(Spacer(1, 2))

    def on_page(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor(th["muted"]))
        canvas.drawString(0.85 * inch, 0.5 * inch, footer[:140])
        canvas.drawRightString(LETTER[0] - 0.85 * inch, 0.5 * inch, f"Page {d.page}")
        canvas.setStrokeColor(colors.HexColor(th["primary"]))
        canvas.setLineWidth(2)
        canvas.line(0.85 * inch, LETTER[1] - 0.55 * inch, LETTER[0] - 0.85 * inch, LETTER[1] - 0.55 * inch)
        canvas.restoreState()

    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return buf.getvalue()


def _add_hyperlink(paragraph, url: str, text: str, color: str):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color.lstrip("#"))
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(c)
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _docx_spans(p, spans, th):
    for s in spans:
        text = clean_text(s.text)
        if s.kind == "link":
            _add_hyperlink(p, s.url, text, th["primary"])
        else:
            parts = text.split("\n")
            for i, part in enumerate(parts):
                run = p.add_run(part)
                run.bold = s.kind == "bold"
                if i < len(parts) - 1:
                    run.add_break()


def render_docx(title: str, subtitle: str, sections: list[ReportSection], footer: str, theme: dict | None = None) -> bytes:
    th = {**THEME, **(theme or {})}
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    core = doc.core_properties
    core.title, core.author, core.subject = title, "EcoVG PitchForce", subtitle
    core.created = datetime.utcnow()
    h = doc.add_heading(title, level=0)
    for r in h.runs:
        r.font.color.rgb = RGBColor.from_string(th["primary"].lstrip("#"))
    p = doc.add_paragraph(subtitle)
    p.runs[0].font.color.rgb = RGBColor.from_string(th["muted"].lstrip("#"))
    for sec in sections:
        if not sec.text.strip():
            continue
        if sec.key != "header":
            hh = doc.add_heading(sec.title, level=1)
            for r in hh.runs:
                r.font.color.rgb = RGBColor.from_string(th["primary"].lstrip("#"))
        for b in markup.parse(sec.text):
            if b.kind == "heading":
                hp = doc.add_heading(level=2)
                _docx_spans(hp, b.spans, th)
            elif b.kind == "bullets":
                for item in b.items:
                    bp = doc.add_paragraph(style="List Bullet")
                    _docx_spans(bp, item, th)
            else:
                para = doc.add_paragraph()
                _docx_spans(para, b.spans, th)
                if b.kind == "note":
                    for r in para.runs:
                        r.italic = True
    fp = doc.sections[0].footer.paragraphs[0]
    fp.text = footer
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
