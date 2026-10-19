"""Exports open and are readable; founder packages contain only their own company's information."""
import io
import json
import zipfile

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from pitchforce.models import Artifact, Company, Event
from pitchforce.reports.package import build_company_package, build_event_package
from pitchforce.storage import get_storage

from .conftest import company

OTHER_NAMES = ["Brightwell", "Owlvision", "Corvane", "Hemaflow", "Ana Silva", "Ravi Mehta", "Sam Whitfield", "Dana Okafor", "Jamie Fox",
               "Lee Grant", "Pat Rivera"]


def _zip_text(data: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return {n: z.read(n) for n in z.namelist()}


def test_company_package_contents_and_isolation(fresh_sample):
    env, eid = fresh_sample
    db = env.session()
    ev = db.get(Event, eid)
    hema = company(db, eid, "Hemaflow Bio")
    res = build_company_package(db, ev, hema)
    db.commit()
    z = db.get(Artifact, res["zip"])
    files = _zip_text(get_storage().read_bytes(z.storage_key))
    names = sorted(n.split("/", 1)[1] for n in files)
    assert names == sorted(["Hemaflow_Bio_PitchForce_Feedback.pdf", "Hemaflow_Bio_PitchForce_Feedback.docx",
                            "Hemaflow_Bio_Investor_Follow-up.xlsx", "Hemaflow_Bio_email_draft.txt",
                            "Hemaflow_Bio_Evidence_Appendix.html", "Hemaflow_Bio_manifest.json"])
    pdf = next(v for k, v in files.items() if k.endswith(".pdf"))
    text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    assert "Hemaflow Bio" in text and "Top three improvement priorities" in text and 1 <= len(PdfReader(io.BytesIO(pdf)).pages) <= 4
    docx = Document(io.BytesIO(next(v for k, v in files.items() if k.endswith(".docx"))))
    dtext = "\n".join(p.text for p in docx.paragraphs)
    assert "Overall assessment" in dtext and "Announced total: 21 / 30" in dtext
    wb = load_workbook(io.BytesIO(next(v for k, v in files.items() if k.endswith(".xlsx"))))
    assert wb.sheetnames == ["Follow-up", "Event responses", "Summary"]
    fu = [[c.value for c in row] for row in wb["Follow-up"].iter_rows(min_row=2)]
    assert len(fu) == 1 and fu[0][1] == "Pat Rivera" and fu[0][0] == "Information request"  # not-interested rows excluded
    resp = [[c.value for c in row] for row in wb["Event responses"].iter_rows(min_row=2)]
    assert len(resp) == 5  # only the five event-specific responses; historical rows excluded
    manifest = json.loads(next(v for k, v in files.items() if k.endswith("manifest.json")))
    assert any("Video clip not included" in x for x in manifest["limitations"])
    email = next(v for k, v in files.items() if k.endswith(".txt")).decode()
    assert "Hi Dana" in email and "asked for more information" in email and "interest in investing" not in email
    assert "video clip" not in email.lower()  # no clip was generated, so the email must not mention one
    everything = b"\n".join(files.values()).decode("latin-1")
    for other in OTHER_NAMES:
        if other in ("Hemaflow", "Dana Okafor", "Pat Rivera"):
            continue
        assert other not in everything, other
    db.close()


def test_event_package_has_organizer_files(fresh_sample):
    env, eid = fresh_sample
    db = env.session()
    ev = db.get(Event, eid)
    res = build_event_package(db, ev, include_clips=False)
    db.commit()
    assert all(v == "ok" for v in res["results"].values())
    z = db.get(Artifact, res["zip"])
    files = _zip_text(get_storage().read_bytes(z.storage_key))
    assert "organizer/Event_Audit_Workbook.xlsx" in files and "organizer/structured_export.json" in files
    assert len([n for n in files if n.startswith("companies/")]) == 4
    wb = load_workbook(io.BytesIO(files["organizer/Event_Audit_Workbook.xlsx"]))
    assert {"Results", "Judge scores", "Poll launches", "All poll responses", "Review items", "Sources"} <= set(wb.sheetnames)
    results = {r[0].value: r for r in wb["Results"].iter_rows(min_row=2)}
    assert results["Brightwell Diagnostics"][8].value == 18 and results["Brightwell Diagnostics"][7].value == 3  # imputed shown
    export = json.loads(files["organizer/structured_export.json"])
    assert export["schema"].startswith("pitchforce-feedback/") and len(export["companies"]) == 4
    db.close()


def test_rebuild_is_idempotent(fresh_sample):
    env, eid = fresh_sample
    db = env.session()
    ev = db.get(Event, eid)
    c = company(db, eid, "Corvane Pharma")
    build_company_package(db, ev, c)
    build_company_package(db, ev, c)
    db.commit()
    current = db.query(Artifact).filter_by(company_id=c.id, is_current=True).all()
    kinds = [a.kind for a in current]
    assert len(kinds) == len(set(kinds))  # one current file per kind, no duplicates
    db.close()


_ = Company
