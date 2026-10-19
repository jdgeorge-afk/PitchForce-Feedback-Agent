import io
from datetime import date

from openpyxl import Workbook, load_workbook

from pitchforce.models import Company, ContactMatch, Event, Person, PollLaunch, PollResponse, ReviewItem
from pitchforce.parsing import polls as P
from pitchforce.parsing.rosters import parse_company_roster, parse_judge_roster
from pitchforce.parsing.tabular import read_tabular
from pitchforce.sample_data import COMPANY_ROSTER, poll_csv_hemaflow, write_csv
from pitchforce.services import interest
from pitchforce.services.ingest import ingest_source
from pitchforce.services.sources import register_bytes

from .conftest import company


def test_standard_csv_and_xlsx_roster_import():
    rows = parse_company_roster(read_tabular(write_csv(COMPANY_ROSTER), "r.csv")).rows
    assert [r.values["company"] for r in rows][:2] == ["Brightwell Diagnostics", "Owlvision"]
    assert rows[1].values["presenter"] == "Sam Whitfield" and rows[1].original["Phone"] == "+1 555 0100"
    wb = Workbook()
    ws = wb.active
    ws.append(["First Name", "Last Name", "Firm", "Email", "Attendance"])
    ws.append(["Ada", "Stone", "Stone Capital", "", "Absent"])
    buf = io.BytesIO()
    wb.save(buf)
    j = parse_judge_roster(read_tabular(buf.getvalue(), "j.xlsx")).rows
    assert j[0].values["name"] == "Ada Stone" and j[0].values["email"] == "" and j[0].values["attendance"] == "Absent"


def test_zoom_multisection_csv_with_quoted_newlines_and_trailing_columns():
    parsed = P.parse_poll_sheets(read_tabular(poll_csv_hemaflow(), "report.csv"))
    assert parsed.layout == "zoom_multi_section"
    assert len(parsed.launches) == 4 and parsed.launches[2].meeting_id_raw == "812 3456 7890"
    assert len(parsed.responses) == 10
    multi = next(r for r in parsed.responses if r.rid == "1")
    assert P.split_choices(multi.answer) == ["I'm interested in investing.", "I'd like a meeting."]  # quoted multiline kept
    assert all("(blank column" not in k for r in parsed.responses for k in r.original if not r.original[k])


def test_categories_unicode_and_distinct():
    assert P.categorize("I’m not interested.") == "not_interested"  # curly apostrophe
    assert P.categorize("I need more info.") == "more_info"
    assert P.categorize("I'm interested in investing.") == "invest"
    assert P.categorize("I have a free resource.") == "free_resource"
    assert P.categorize("Something else entirely") == "unknown"
    assert P.is_contradictory({"not_interested", "meeting"}) and not P.is_contradictory({"more_info"})


def _event(db, **kw):
    ev = Event(name="Synthetic", event_date=date(2026, 9, 17), timezone="America/Los_Angeles", **kw)
    db.add(ev)
    db.flush()
    return ev


def _load(db, ev):
    ingest_source(db, register_bytes(db, ev, "company_roster", "roster.csv", write_csv(COMPANY_ROSTER)))
    judges = write_csv([["Name", "Email"], ["Tom Becker", "tom.becker@example.org"]])
    ingest_source(db, register_bytes(db, ev, "judge_roster", "judges.csv", judges))
    src = register_bytes(db, ev, "interest", "report_Follow-Up- Presenting Hemaflow Bio.csv", poll_csv_hemaflow())
    ingest_source(db, src)
    db.flush()
    return src


def test_event_filtering_excludes_historical_and_counts_people_once(db):
    ev = _event(db, webinar_id="812 3456 7890")
    _load(db, ev)
    launches = db.query(PollLaunch).all()
    assert [pl.in_event for pl in sorted(launches, key=lambda x: x.launch_index)] == [False, False, True, False]
    in_ev = db.query(PollResponse).filter_by(in_event=True).all()
    assert sorted(r.original_response_id for r in in_ev) == ["49", "50", "51", "52", "53"]
    hema = company(db, ev.id, "Hemaflow Bio")
    s = interest.summarize(db, ev, hema)
    assert s.status == "available" and s.respondents == 5 and s.anonymous == 1
    assert s.by_category == {"more_info": 1, "not_interested": 4}
    assert "invest" not in s.by_category  # "needs more info" is not investment interest


def test_launch_company_needs_more_than_filename(db):
    ev = _event(db)
    _load(db, ev)
    pl = db.query(PollLaunch).filter_by(in_event=True).one()
    kinds = {e["kind"] for e in pl.company_evidence["items"]}
    assert "poll name" in kinds  # strong evidence from the poll name, not only the filename
    assert pl.company_status == "auto"  # the poll name names the company (the filename only agrees)
    # Filename-only evidence must not be enough:
    ev2 = _event(db)
    ingest_source(db, register_bytes(db, ev2, "company_roster", "r2.csv", write_csv(COMPANY_ROSTER)))
    csv2 = poll_csv_hemaflow().replace(b"Follow-Up- Presenting Hemaflow Bio", b"Follow-Up poll")
    src2 = register_bytes(db, ev2, "interest", "report_Follow-Up- Presenting Hemaflow Bio v2.csv", csv2)
    ingest_source(db, src2)
    pl2 = db.query(PollLaunch).filter_by(event_id=ev2.id, in_event=True).one()
    assert pl2.company_status in ("proposed", "unassigned") and pl2.company_status != "auto"
    assert db.query(ReviewItem).filter_by(event_id=ev2.id, kind="interest_assignment").count() >= 1
    hema2 = company(db, ev2.id, "Hemaflow Bio")
    assert interest.summarize(db, ev2, hema2).status == "needs_review"


def test_anonymous_and_conflicting_contacts(db):
    ev = _event(db)
    _load(db, ev)
    anon = db.query(PollResponse).filter_by(original_response_id="52").one()
    assert anon.is_anonymous and anon.name == "" and anon.email == ""  # no invented contact details
    r50 = db.query(PollResponse).filter_by(original_response_id="50").one()
    assert r50.name == "Alex Marin" and r50.email == "tom.becker@example.org"
    cm = db.query(ContactMatch).filter_by(response_id=r50.id).one()
    assert cm.status == "conflict"
    tom = db.query(Person).filter_by(display_name="Tom Becker").one()
    assert tom.display_name == "Tom Becker" and not tom.aliases  # not merged
    assert db.query(ReviewItem).filter(ReviewItem.dedupe_key.like("email-conflict:%")).count() == 1


def test_missing_interest_differs_from_zero(db):
    ev = _event(db)
    ingest_source(db, register_bytes(db, ev, "company_roster", "roster.csv", write_csv(COMPANY_ROSTER)))
    hema = company(db, ev.id, "Hemaflow Bio")
    assert interest.summarize(db, ev, hema).status == "missing_source"
    _load(db, ev)
    corvane = company(db, ev.id, "Corvane Pharma")
    s = interest.summarize(db, ev, corvane)
    assert s.status == "not_assigned" and "not zero" in s.note


def test_duplicates_amendments_and_contradictions(db):
    from pitchforce.sample_data import interest_xlsx_other

    ev = _event(db)
    ingest_source(db, register_bytes(db, ev, "company_roster", "roster.csv", write_csv(COMPANY_ROSTER)))
    ingest_source(db, register_bytes(db, ev, "interest", "other.xlsx", interest_xlsx_other()))
    bw = company(db, ev.id, "Brightwell Diagnostics")
    rows = db.query(PollResponse).filter_by(company_id=bw.id, in_event=True).all()
    jamie = sorted([r for r in rows if r.name == "Jamie Fox"], key=lambda r: r.submitted_at)
    assert jamie[0].superseded_by_id == jamie[1].id  # amended submission supersedes, original kept
    lee = next(r for r in rows if r.name == "Lee Grant")
    assert lee.has_conflict and {c.category for c in lee.choices} == {"not_interested", "meeting"}
    s = interest.summarize(db, ev, bw)
    assert s.respondents == 3  # Jamie counted once despite two submissions and three selections
    assert s.by_category["invest"] == 1 and s.by_category["introduction"] == 1 and s.by_category["meeting"] == 2
    # Unknown company label in a row -> review, not silently assigned
    assert db.query(ReviewItem).filter(ReviewItem.dedupe_key.like("row-company:%")).count() == 1


def test_founder_workbook_formula_injection_and_scope(db):
    from pitchforce.reports.xlsx import founder_workbook
    from pitchforce.sample_data import interest_xlsx_other
    from pitchforce.services import matching

    ev = _event(db)
    ingest_source(db, register_bytes(db, ev, "company_roster", "roster.csv", write_csv(COMPANY_ROSTER)))
    src = register_bytes(db, ev, "interest", "other.xlsx", interest_xlsx_other())
    owl = company(db, ev.id, "Owlvision")
    matching.add_company_alias(db, owl, "SightCaption", "operator", status="confirmed")
    ingest_source(db, src)
    data, meta = founder_workbook(db, ev, owl)
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Follow-up"]
    names = [c.value for c in ws["B"]][1:]
    assert any(isinstance(n, str) and n.startswith("'=HYPERLINK") for n in names)  # neutralized, not a formula
    assert all(c.data_type != "f" for row in ws.iter_rows() for c in row)
    assert "Jamie Fox" not in names  # another company's contact is never included
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref
    sub = ws.cell(row=2, column=12).value
    assert sub.__class__.__name__ == "datetime"  # typed dates
